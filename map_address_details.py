"""Replace address IDs in the 0929 workbook using input3's mapping CSV."""
import argparse
import csv
from datetime import datetime
import os
from pathlib import Path
import shutil
import sys

from fill_addresses import temporary_path

ROOT = Path(__file__).resolve().parent
WORKBOOK_NAME = 'Transaction Columns for Fintech RFI_0929.xlsx'
CSV_NAME = 'address_mapping_to_Alert.csv'
TARGETS = [f'{party} Address Line 1' for party in
           ('Debtor', 'Creditor', 'Initiating Party')]


def read_mapping(path, encoding='utf-8-sig'):
    result = {}
    with path.open(encoding=encoding, newline='') as handle:
        reader = csv.DictReader(handle, strict=True)
        fields = reader.fieldnames
        if fields is None:
            raise ValueError('Mapping CSV has no header.')
        reader.fieldnames = [name.strip() for name in fields]
        if len(set(reader.fieldnames)) != len(fields):
            raise ValueError('Mapping CSV has duplicate headers.')
        if not {'raw_addr_id', 'raw_addr_details'} <= set(reader.fieldnames):
            raise ValueError('Mapping CSV requires raw_addr_id and raw_addr_details.')
        for record in reader:
            if None in record or any(value is None for value in record.values()):
                raise ValueError(f'CSV line {reader.line_num}: incorrect field count.')
            key = record['raw_addr_id'].strip()
            if not key:
                continue
            details = record['raw_addr_details']
            if len(details) > 32767:
                raise ValueError(f'CSV line {reader.line_num}: address exceeds Excel cell limit.')
            if key in result and result[key] != details:
                raise ValueError(f'CSV line {reader.line_num}: conflicting addresses for the same ID.')
            result[key] = details
    return result


def update_workbook(path, mapping, sheet_name=None):
    from openpyxl import load_workbook
    workbook = load_workbook(path)
    temporary = None
    try:
        def headers(sheet):
            return [cell.value.strip() if isinstance(cell.value, str) else cell.value
                    for cell in sheet[1]]

        candidates = [sheet for sheet in workbook if set(TARGETS) <= set(headers(sheet))]
        if sheet_name:
            sheet = workbook[sheet_name]
        elif len(candidates) == 1:
            sheet = candidates[0]
        else:
            raise ValueError('Cannot identify a unique worksheet; specify --sheet NAME.')
        labels = headers(sheet)
        if any(labels.count(name) != 1 for name in TARGETS):
            raise ValueError('Worksheet requires each of the three Address Line 1 headers exactly once.')
        stats = {name: dict(matched=0, unmatched=0, blank=0) for name in TARGETS}
        for name in TARGETS:
            column = labels.index(name) + 1
            for row in range(2, sheet.max_row + 1):
                cell = sheet.cell(row, column)
                value = cell.value
                if value is None or isinstance(value, str) and not value.strip():
                    stats[name]['blank'] += 1
                    continue
                if cell.data_type == 'f' or not isinstance(value, str):
                    raise ValueError(f'{sheet.title}!{cell.coordinate}: address ID must be text, not a number or formula.')
                key = value.strip()
                if key not in mapping:
                    stats[name]['unmatched'] += 1
                    continue
                cell.value = mapping[key] or None
                if cell.value is not None:
                    cell.data_type = 's'
                stats[name]['matched'] += 1
        temporary = temporary_path(path.parent, '.xlsx')
        workbook.save(temporary)
        backup = path.with_name(f'{path.stem}.backup_{datetime.now():%Y%m%d_%H%M%S_%f}.xlsx')
        shutil.copy2(path, backup)
        workbook.close()
        os.replace(temporary, path)
        return stats, backup
    finally:
        workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input3', type=Path, default=ROOT / 'input3')
    parser.add_argument('--encoding', default='utf-8-sig')
    parser.add_argument('--sheet', help='Worksheet name; otherwise detected from row 1 headers.')
    args = parser.parse_args()
    try:
        mapping = read_mapping(args.input3 / CSV_NAME, args.encoding)
        stats, backup = update_workbook(args.input3 / WORKBOOK_NAME, mapping, args.sheet)
        for name, counts in stats.items():
            print(f'{name}: matched={counts["matched"]}, unmatched (kept)={counts["unmatched"]}, blank={counts["blank"]}')
        print(f'Updated: {args.input3 / WORKBOOK_NAME}\nBackup: {backup}')
    except ImportError:
        print('Install dependencies: py -3 -m pip install -r requirements.txt', file=sys.stderr)
        return 1
    except (ValueError, OSError, UnicodeError, csv.Error, KeyError) as error:
        print(f'Address mapping failed: {error}', file=sys.stderr)
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
