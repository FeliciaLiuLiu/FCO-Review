#!/usr/bin/env python3
"""Combine eight address batches, then update only 15 address columns. Python 3.11."""
import argparse
import csv
from datetime import datetime
import os
from pathlib import Path
import shutil
import sys
import tempfile

ROOT = Path(__file__).resolve().parent
WORKBOOK_NAME = 'Transaction Columns for Fintech RFI_0928.xlsx'
TRANSACTION_KEY = 'Transaction ID (UETR or Other Unique ID)'
BATCH_KEY = 'cinq_trxn_id'
ADDRESS_MAPPING = [
    (f'{party} Address Line {line}', f'{prefix}{line}')
    for party, prefix in [
        ('Ultimate Debtor', 'ult_deb_address_line'),
        ('Debtor', 'orig_addr_line'),
        ('Creditor', 'secondary_bene_addr_line'),
        ('Ultimate Creditor', 'ult_cred_address_line'),
        ('Initiating Party', 'ip_address_line'),
    ]
    for line in (1, 2, 3)
]

def checked_headers(fields, required, location):
    if not fields or any(not name for name in fields) or len(fields) != len(set(fields)):
        raise ValueError(f'{location}: 表头为空、存在空列名或重复列名')
    missing = sorted(set(required) - set(fields))
    if missing:
        raise ValueError(f'{location}: 缺少列: {", ".join(missing)}')

def normalize_key(value):
    if value is None:
        return ''
    if not isinstance(value, str):
        raise ValueError('交易 ID 必须为文本；请确认 Excel 中的 ID 没有被转为数字而丢失前导零或精度')
    return value.strip()

def temporary_path(folder, suffix):
    descriptor, name = tempfile.mkstemp(dir=folder, suffix=suffix)
    os.close(descriptor)
    return Path(name)

def combine_batches(folder, encoding):
    paths = [folder / f'Batch{i}.csv' for i in range(1, 9)]
    missing = [p.name for p in paths if not p.is_file()]
    if missing:
        raise ValueError(f'{folder}: 缺少文件: {", ".join(missing)}')
    required = [BATCH_KEY] + [source for _, source in ADDRESS_MAPPING]
    headers = []
    # Preserve all source columns, including extra columns unique to a batch.
    for path in paths:
        with path.open(encoding=encoding, newline='') as handle:
            fields = csv.DictReader(handle, strict=True).fieldnames
            checked_headers(fields, required, path.name)
            headers.extend(name for name in fields if name not in headers)
    temporary = temporary_path(folder, '.csv')
    counts = []
    try:
        with temporary.open('w', encoding='utf-8-sig', newline='') as output:
            writer = csv.DictWriter(output, fieldnames=headers)
            writer.writeheader()
            for path in paths:
                count = 0
                with path.open(encoding=encoding, newline='') as handle:
                    reader = csv.DictReader(handle, strict=True)
                    for record in reader:
                        if None in record or any(value is None for value in record.values()):
                            raise ValueError(f'{path.name}, 行 {reader.line_num}: 字段数量与表头不一致')
                        writer.writerow(record)
                        count += 1
                counts.append((path.name, count))
        destination = folder / 'batch_combined.csv'
        os.replace(temporary, destination)
        return destination, counts
    finally:
        temporary.unlink(missing_ok=True)

def address_index(combined):
    result = {}
    duplicates = 0
    blank_keys = 0
    with combined.open(encoding='utf-8-sig', newline='') as handle:
        reader = csv.DictReader(handle, strict=True)
        for record in reader:
            key = normalize_key(record[BATCH_KEY])
            if not key:
                blank_keys += 1
                continue
            addresses = tuple(record[source] if record[source].strip() else None
                              for _, source in ADDRESS_MAPPING)
            if key in result:
                if result[key] != addresses:
                    raise ValueError(f'batch_combined.csv 行 {reader.line_num}: 同一 cinq_trxn_id 存在不同地址；请修正后重试，Excel 未更新')
                duplicates += 1
            else:
                result[key] = addresses
    return result, duplicates, blank_keys

def update_workbook(path, index, sheet_name=None):
    from openpyxl import load_workbook
    workbook = load_workbook(path)
    temporary = None
    try:
        if sheet_name:
            sheet = workbook[sheet_name]
        else:
            candidates = [sheet for sheet in workbook if TRANSACTION_KEY in
                          [cell.value for cell in sheet[1]]]
            if len(candidates) != 1:
                raise ValueError('无法唯一确定交易工作表，请使用 --sheet 指定工作表名')
            sheet = candidates[0]
        headers = [cell.value for cell in sheet[1]]
        required = [TRANSACTION_KEY] + [target for target, _ in ADDRESS_MAPPING]
        for label in required:
            if headers.count(label) != 1:
                raise ValueError(f'Excel 必需列缺失或重复: {label}')
        key_column = headers.index(TRANSACTION_KEY) + 1
        address_columns = [headers.index(target) + 1 for target, _ in ADDRESS_MAPPING]
        matched = unmatched = blank_keys = 0
        for row in sheet.iter_rows(min_row=2):
            if all(cell.value is None for cell in row):
                continue
            key_cell = row[key_column - 1]
            if key_cell.data_type == 'f':
                raise ValueError(f'Excel 行 {key_cell.row}: 交易 ID 是公式，请改为原始文本 ID')
            key = normalize_key(key_cell.value)
            if key and key in index:
                addresses = index[key]
                matched += 1
            else:
                addresses = (None,) * 15
                unmatched += 1
                blank_keys += int(not key)
            for column, value in zip(address_columns, addresses):
                cell = row[column - 1]
                cell.value = value
                if isinstance(value, str):
                    cell.data_type = 's'  # Literal text even when an address starts with '='.
        temporary = temporary_path(path.parent, '.xlsx')
        workbook.save(temporary)
        # Verify saved addresses before replacing the user's workbook.
        check = load_workbook(temporary, read_only=True)
        try:
            for row in check[sheet.title].iter_rows(min_row=2, values_only=True):
                if all(value is None for value in row):
                    continue
                expected = index.get(normalize_key(row[key_column - 1]), (None,) * 15)
                actual = tuple(row[column - 1] for column in address_columns)
                if actual != expected:
                    raise ValueError('导出后的地址校验失败；原 Excel 未更新')
        finally:
            check.close()
        backup = path.with_name(f'{path.stem}.backup_{datetime.now():%Y%m%d_%H%M%S_%f}.xlsx')
        shutil.copy2(path, backup)
        workbook.close()
        os.replace(temporary, path)
        return matched, unmatched, blank_keys, backup
    finally:
        workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)

def main():
    parser = argparse.ArgumentParser(description='合并 input2 的 8 个 Batch CSV，并补充 Excel 的 15 列地址')
    parser.add_argument('--input2', type=Path, default=ROOT / 'input2')
    parser.add_argument('--workbook', type=Path, default=ROOT / 'output' / WORKBOOK_NAME)
    parser.add_argument('--encoding', default='utf-8-sig')
    parser.add_argument('--sheet', help='默认自动识别包含交易 ID 表头的唯一工作表')
    args = parser.parse_args()
    try:
        combined, counts = combine_batches(args.input2, args.encoding)
        for name, count in counts:
            print(f'{name}: {count} 行')
        print(f'合并完成: {combined}，共 {sum(count for _, count in counts)} 行')
        index, duplicates, blank = address_index(combined)
        print(f'唯一交易 ID: {len(index)}；相同地址重复 ID: {duplicates}；空 ID: {blank}')
        if not args.workbook.is_file():
            raise ValueError(f'找不到 Excel: {args.workbook}；合并 CSV 已生成，可用 --workbook 指定路径')
        matched, unmatched, empty, backup = update_workbook(args.workbook, index, args.sheet)
        print(f'完成：匹配 {matched} 行；未匹配 {unmatched} 行（含空交易 ID {empty} 行），未匹配地址留空')
        print(f'Excel: {args.workbook}\n原文件备份: {backup}')
    except ImportError:
        print('请先运行: python -m pip install -r requirements.txt', file=sys.stderr)
        return 1
    except (ValueError, OSError, UnicodeError, csv.Error, KeyError) as error:
        print(f'地址补充失败: {error}', file=sys.stderr)
        return 1
    return 0

if __name__ == '__main__':
    sys.exit(main())
