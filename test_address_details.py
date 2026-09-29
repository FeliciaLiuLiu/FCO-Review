"""Temporary fixtures verify mapping and preservation of unrelated workbook data."""
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from map_address_details import TARGETS, read_mapping, update_workbook


class AddressDetailsTests(unittest.TestCase):
    def test_mapping_and_preservation(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'test.xlsx'
            csv_path = Path(folder) / 'mapping.csv'
            csv_path.write_text('raw_addr_id,raw_addr_details\n001,"Street, City"\n002,=literal\n003,\n', encoding='utf-8-sig')
            book = Workbook()
            sheet = book.active
            sheet.append([' ' + TARGETS[0], *TARGETS[1:], 'Debtor Address Line 2', 'Other'])
            sheet.append(['001', '002', '003', 'keep line 2', '=1+1'])
            sheet.append(['missing', None, '001', 'keep too', 42])
            sheet['A2'].fill = PatternFill('solid', fgColor='FFFF00')
            book.create_sheet('Other sheet')['A1'] = 'unchanged'
            book.save(path)
            book.close()
            original = path.read_bytes()
            stats, backup = update_workbook(path, read_mapping(csv_path))
            self.assertEqual(backup.read_bytes(), original)
            result = load_workbook(path)
            try:
                sheet = result.active
                self.assertEqual([cell.value for cell in sheet[2]],
                                 ['Street, City', '=literal', None, 'keep line 2', '=1+1'])
                self.assertEqual(sheet['B2'].data_type, 's')
                self.assertEqual(sheet['E2'].data_type, 'f')
                self.assertEqual(sheet['A2'].fill.fgColor.rgb, '00FFFF00')
                self.assertEqual(sheet['A3'].value, 'missing')
                self.assertEqual(sheet['C3'].value, 'Street, City')
                self.assertEqual(result['Other sheet']['A1'].value, 'unchanged')
                self.assertEqual(stats[TARGETS[0]]['unmatched'], 1)
            finally:
                result.close()

    def test_conflicting_mapping(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'mapping.csv'
            path.write_text('raw_addr_id,raw_addr_details\n001,A\n001,B\n', encoding='utf-8')
            with self.assertRaisesRegex(ValueError, 'conflicting'):
                read_mapping(path)

    def test_invalid_id_keeps_original(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'test.xlsx'
            book = Workbook()
            book.active.append(TARGETS)
            book.active.append(['001', 123, '001'])
            book.save(path)
            book.close()
            original = path.read_bytes()
            with self.assertRaisesRegex(ValueError, 'must be text'):
                update_workbook(path, {'001': 'Address'})
            self.assertEqual(path.read_bytes(), original)
            self.assertEqual(list(Path(folder).iterdir()), [path])


if __name__ == '__main__':
    unittest.main()
