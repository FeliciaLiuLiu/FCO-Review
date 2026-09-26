#!/usr/bin/env python3
"""Combine customer CSVs into the requested Excel workbook. Python 3.9+."""
import argparse
import csv
import math
from decimal import Decimal, InvalidOperation, localcontext
from pathlib import Path
import re
import sys
import os
import tempfile

from mapping import MAPPING, HEADERS

ROOT = Path(__file__).resolve().parent
DEFAULT_OUTPUT = ROOT / 'output' / 'Transaction Columns for Fintech RFI_0925_v1.xlsx'

def export_excel(rows, output):
    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ValueError('请先运行：python -m pip install -r requirements.txt') from None
    book = Workbook(write_only=True)
    sheet = book.create_sheet('Transactions')
    sheet.freeze_panes = 'B2'
    sheet.sheet_view.showGridLines = False
    for index in range(1, 46):
        sheet.column_dimensions[get_column_letter(index)].width = 36 if index == 3 else 25
    sheet.row_dimensions[1].height = 64
    header_cells = []
    for label in HEADERS:
        cell = WriteOnlyCell(sheet, value=label)
        cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='203864')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        header_cells.append(cell)
    sheet.append(header_cells)
    for row in rows:
        cells = []
        for index, value in enumerate(row):
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.data_type = 's'
                cell.number_format = '@'
            elif index == 3:
                cell.number_format = '0.00#############'
            cell.font = Font(name='Arial', size=10)
            cells.append(cell)
        sheet.append(cells)
    output.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(suffix='.xlsx', dir=output.parent)
    os.close(descriptor)
    try:
        book.save(temporary)
        os.replace(temporary, output)
    finally:
        Path(temporary).unlink(missing_ok=True)

def amount_value(raw, location):
    if not raw.strip():
        return None
    if not re.fullmatch(r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?', raw.strip()):
        raise ValueError(f'{location}: amount 不是有效数字')
    try:
        with localcontext() as context:
            context.prec = max(28, len(raw) + 4)
            value = Decimal(raw.strip()) / Decimal(100)
        number = float(value)
        if not math.isfinite(number) or Decimal(str(number)) != value:
            raise ValueError(f'{location}: amount 超出 Excel 可安全表示的数值范围')
        if len(value.normalize().as_tuple().digits) > 15:
            raise ValueError(f'{location}: amount 超过 Excel 的 15 位有效数字精度')
        return number
    except (InvalidOperation, OverflowError):
        raise ValueError(f'{location}: amount 不是有效数字') from None

def read_rows(folder, encoding):
    files = sorted((p for p in folder.iterdir() if p.is_file() and p.suffix.lower() == '.csv'),
                   key=lambda p: p.name)
    if not files:
        raise ValueError('input 文件夹没有 CSV；如需空白模板，请使用 --empty-template')
    required = {source for _, source in MAPPING if source and source != '$filename'}
    rows = []
    counts = []
    for path in files:
        count = 0
        with path.open('r', encoding=encoding, newline='') as handle:
            reader = csv.DictReader(handle, strict=True)
            fields = reader.fieldnames
            if not fields or len(fields) != len(set(fields)):
                raise ValueError(f'{path.name}: CSV 表头为空或存在重复列名')
            missing = sorted(required - set(fields))
            if missing:
                raise ValueError(f'{path.name}: 缺少列: {", ".join(missing)}')
            for record in reader:
                location = f'{path.name}, CSV 行 {reader.line_num}'
                if None in record or any(v is None for v in record.values()):
                    raise ValueError(f'{location}: 字段数量与表头不一致')
                output = []
                for _, source in MAPPING:
                    if source == '$filename':
                        value = path.stem
                    elif source is None:
                        value = None
                    elif source == 'amount':
                        value = amount_value(record[source], location)
                    else:
                        value = record[source] if record[source].strip() else None
                    output.append(value)
                rows.append(output)
                count += 1
                if len(rows) > 1048575:
                    raise ValueError('数据超过单个 Excel 工作表允许的 1,048,575 条数据行')
        counts.append((path.name, count))
    return rows, counts

def main():
    parser = argparse.ArgumentParser(description='将 input 中客户 CSV 合并为 45 列 Excel')
    parser.add_argument('--input', type=Path, default=ROOT / 'input')
    parser.add_argument('--output', type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument('--encoding', default='utf-8-sig', help='默认 UTF-8，兼容 BOM；可指定 gb18030')
    parser.add_argument('--empty-template', action='store_true', help='忽略 CSV，仅生成表头')
    args = parser.parse_args()
    try:
        rows, counts = ([], []) if args.empty_template else read_rows(args.input, args.encoding)
        export_excel(rows, args.output.resolve())
        for name, count in counts:
            print(f'{name}: {count} 行')
        print(f'完成：{len(rows)} 条数据，{len(HEADERS)} 列。输出：{args.output.resolve()}')
    except (ValueError, OSError, UnicodeError, csv.Error) as error:
        print(f'转换失败：{error}', file=sys.stderr)
        return 1
    return 0

if __name__ == '__main__':
    sys.exit(main())
